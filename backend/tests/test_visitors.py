from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import visitors
from app.main import app


client = TestClient(app)


def test_visitor_rows_are_appended(tmp_path, monkeypatch) -> None:
    log = tmp_path / "visitors.xlsx"
    monkeypatch.setattr(visitors, "VISITOR_LOG", log)
    entry = {"name": "Asha Rao", "email": "asha@example.com", "organisation": "Acme", "role": "Planner", "purpose": "Demo"}
    assert client.post("/api/visitors", json=entry).status_code == 200
    assert client.post("/api/visitors", json={**entry, "visit": "return"}).status_code == 200
    rows = list(load_workbook(log)["Visitors"].iter_rows(values_only=True))
    assert rows[0] == tuple(visitors.COLUMNS)
    assert [row[1] for row in rows[1:]] == ["Asha Rao", "Asha Rao"]
    assert [row[6] for row in rows[1:]] == ["first", "return"]


def test_visitor_needs_a_valid_email(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(visitors, "VISITOR_LOG", tmp_path / "visitors.xlsx")
    response = client.post("/api/visitors", json={"name": "Asha", "email": "not-an-email"})
    assert response.status_code == 422
