"""Visitor sign-in log: one row per visit, appended to an Excel workbook on the server."""

import os
import threading
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from .config import BACKEND_DIR

VISITOR_LOG = Path(os.environ.get("DISPATCH_TWIN_VISITOR_LOG", BACKEND_DIR / "data" / "visitors.xlsx"))
COLUMNS = ["Timestamp (UTC)", "Name", "Email", "Organisation", "Role", "Purpose", "Visit", "IP address", "Browser"]

_lock = threading.Lock()


def _open_workbook(path: Path) -> Workbook:
    if path.is_file():
        return load_workbook(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Visitors"
    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for letter, width in zip("ABCDEFGHI", [22, 24, 30, 24, 20, 40, 10, 16, 50]):
        sheet.column_dimensions[letter].width = width
    return workbook


def append_visit(entry: dict, ip: str | None, user_agent: str | None, path: Path | None = None) -> dict:
    path = path or VISITOR_LOG
    timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
    row = [
        timestamp,
        entry["name"],
        entry["email"],
        entry.get("organisation") or "",
        entry.get("role") or "",
        entry.get("purpose") or "",
        entry.get("visit", "first"),
        ip or "",
        (user_agent or "")[:300],
    ]
    with _lock:
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = _open_workbook(path)
        workbook["Visitors"].append(row)
        workbook.save(path)
    return {"recorded_at": timestamp}
